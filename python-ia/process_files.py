from fastapi import APIRouter, HTTPException
from pydantic import BaseModel
import boto3
import os
import google.generativeai as genai
import psycopg2
import io
import re
import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin, urlparse

from pypdf import PdfReader
from docx import Document as DocxDocument
import openpyxl

router = APIRouter()

class ProcessRequest(BaseModel):
    file_name: str
    gemini_api_key: str

class ProcessKnowledgeRequest(BaseModel):
    document_id: int
    file_name: str
    gemini_api_key: str

class ProcessUrlRequest(BaseModel):
    document_id: int
    url: str
    gemini_api_key: str

class CrawlRequest(BaseModel):
    url: str
    max_pages: int = 30

MINIO_ENDPOINT   = os.getenv("MINIO_ENDPOINT", "minio:9000")
MINIO_ACCESS_KEY = os.getenv("MINIO_ACCESS_KEY", "minioadmin")
MINIO_SECRET_KEY = os.getenv("MINIO_SECRET_KEY", "minioadmin")
MINIO_BUCKET_NAME = os.getenv("MINIO_BUCKET_NAME", "materiais")
MINIO_SECURE     = os.getenv("MINIO_SECURE", "false").lower() == "true"

DB_URL = os.getenv("DB_URL")

s3_client = boto3.client(
    's3',
    endpoint_url=f"http://{MINIO_ENDPOINT}" if not MINIO_SECURE else f"https://{MINIO_ENDPOINT}",
    aws_access_key_id=MINIO_ACCESS_KEY,
    aws_secret_access_key=MINIO_SECRET_KEY,
    region_name='us-east-1'
)

def chunk_text(text: str, min_length: int = 50) -> list[str]:
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    text = re.sub(r'[ \t]{2,}', ' ', text)
    text = re.sub(r'\n{4,}', '\n\n', text)
    chunks = [c.strip() for c in re.split(r'\n\n+', text) if len(c.strip()) >= min_length]
    return chunks

def embed_and_store(chunks: list[str], document_id: int, gemini_api_key: str, replace: bool = True):
    """Embeds chunks and stores them in the DB. If replace=True, deletes existing chunks first."""
    genai.configure(api_key=gemini_api_key)
    conn = psycopg2.connect(DB_URL)
    cur = conn.cursor()

    if replace:
        cur.execute('DELETE FROM knowledge_chunks WHERE "documentId" = %s', (document_id,))

    for c in chunks:
        embed_res = genai.embed_content(
            model="models/gemini-embedding-2-preview",
            content=c,
            task_type="retrieval_document",
            output_dimensionality=768
        )
        embedding = embed_res['embedding']
        vector_str = "[" + ",".join(str(x) for x in embedding) + "]"
        cur.execute("""
            INSERT INTO knowledge_chunks (id, "documentId", content, embedding, "createdAt", "updatedAt")
            VALUES (gen_random_uuid(), %s, %s, %s::vector, NOW(), NOW())
        """, (document_id, c, vector_str))

    conn.commit()
    cur.close()
    conn.close()

@router.post("/process")
def process_file(req: ProcessRequest):
    try:
        return {"status": "success", "message": f"Arquivo {req.file_name} recebido."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/process_knowledge")
def process_knowledge(req: ProcessKnowledgeRequest):
    try:
        response = s3_client.get_object(Bucket=MINIO_BUCKET_NAME, Key=req.file_name)
        file_bytes = response['Body'].read()

        text_content = ""
        lower_name = req.file_name.lower()

        if lower_name.endswith(".md") or lower_name.endswith(".txt") or lower_name.endswith(".csv"):
            text_content = file_bytes.decode("utf-8", errors="ignore")
        elif lower_name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(file_bytes))
            parts = [page.extract_text() for page in reader.pages if page.extract_text()]
            text_content = "\n\n".join(parts)
        elif lower_name.endswith(".docx"):
            doc = DocxDocument(io.BytesIO(file_bytes))
            text_content = "\n\n".join([p.text for p in doc.paragraphs if p.text])
        elif lower_name.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    if row_text:
                        text_content += row_text + "\n"
        else:
            return {"status": "error", "message": "Extensão não suportada"}

        if not text_content.strip():
            return {"status": "error", "message": "Nenhum texto extraído"}

        chunks = chunk_text(text_content)
        if not chunks:
            return {"status": "error", "message": "Texto insuficiente para chunking"}

        embed_and_store(chunks, req.document_id, req.gemini_api_key, replace=True)
        return {"status": "success", "chunks_criados": len(chunks)}

    except Exception as e:
        print(f"Erro crítico processando {req.file_name}: {e}")
        return {"status": "error", "message": str(e)}

def _extract_meta(soup, req_url):
    """Extrai title, og:title, og:description, description e texto da página."""
    title_tag = soup.find('title')
    og_title  = soup.find('meta', attrs={'property': 'og:title'})
    og_desc   = soup.find('meta', attrs={'property': 'og:description'})
    meta_desc = soup.find('meta', attrs={'name': 'description'})

    title = ""
    if og_title and og_title.get('content'):
        title = og_title['content']
    elif title_tag:
        title = title_tag.get_text(strip=True)

    description = ""
    if og_desc and og_desc.get('content'):
        description = og_desc['content']
    elif meta_desc and meta_desc.get('content'):
        description = meta_desc['content']

    for tag in soup(['script', 'style', 'nav', 'footer', 'noscript', 'iframe', 'svg', 'form']):
        tag.decompose()
    page_text = soup.get_text(separator='\n', strip=True)

    return title, description, page_text


def _build_content(req_url, title, description, page_text):
    """Monta o conteúdo final para chunking."""
    parts = [f"Fonte: {req_url}"]
    if title:
        parts.append(f"Título: {title}")
    if description:
        parts.append(f"Descrição: {description}")
    if page_text:
        parts.append(page_text)
    return "\n\n".join(parts)


@router.post("/process_url")
def process_url(req: ProcessUrlRequest):
    """Busca uma URL, extrai o texto, cria embeddings e armazena no banco.

    Resiliente a sites que bloqueiam scrapers (ex: Instagram):
    faz fallback para meta tags e cria pelo menos um chunk descritivo.
    """
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        'Accept-Language': 'pt-BR,pt;q=0.9,en;q=0.8',
    }

    title, description, page_text = "", "", ""
    fetch_warning = ""

    try:
        resp = requests.get(req.url, headers=headers, timeout=15, allow_redirects=True)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, 'html.parser')
        title, description, page_text = _extract_meta(soup, req.url)

        # Detecta páginas de login/bloqueio (conteúdo inútil)
        blocked_signals = ["log in", "login", "sign in", "signin", "faça login", "entre para ver"]
        page_lower = page_text.lower()
        if any(s in page_lower for s in blocked_signals) and len(page_text) < 2000:
            fetch_warning = "Site exigiu login — usando apenas meta informações disponíveis."
            page_text = ""

    except requests.exceptions.HTTPError as e:
        status = e.response.status_code if e.response else "?"
        fetch_warning = f"HTTP {status} ao buscar a URL — usando fallback de meta dados."
        # Tenta extrair meta do corpo de erro se houver
        try:
            soup = BeautifulSoup(e.response.text, 'html.parser')
            title, description, _ = _extract_meta(soup, req.url)
        except Exception:
            pass
    except requests.exceptions.ConnectionError:
        return {"status": "error", "message": "Não foi possível conectar à URL. Verifique se está acessível."}
    except requests.exceptions.Timeout:
        return {"status": "error", "message": "Tempo limite ao acessar a URL."}
    except Exception as e:
        print(f"[process_url] Erro ao buscar {req.url}: {e}")
        fetch_warning = f"Erro ao acessar a URL: {str(e)}"

    # Monta conteúdo e chuncka
    full_text = _build_content(req.url, title, description, page_text)
    chunks = chunk_text(full_text)

    # Fallback: se nenhum chunk gerado, cria um chunk mínimo com o que se sabe
    if not chunks:
        fallback = f"URL: {req.url}"
        if title:
            fallback += f"\nTítulo: {title}"
        if description:
            fallback += f"\nDescrição: {description}"
        if fetch_warning:
            fallback += f"\nObservação: {fetch_warning}"
        chunks = [fallback]

    try:
        embed_and_store(chunks, req.document_id, req.gemini_api_key, replace=True)
    except Exception as e:
        print(f"[process_url] Erro ao vetorizar {req.url}: {e}")
        return {"status": "error", "message": f"Erro ao gerar embeddings: {str(e)}"}

    result = {"status": "success", "chunks_criados": len(chunks)}
    if fetch_warning:
        result["aviso"] = fetch_warning
    return result


@router.post("/crawl_site")
def crawl_site(req: CrawlRequest):
    """Rastreia uma URL e retorna todos os links internos do mesmo domínio."""
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        'Accept-Language': 'pt-BR,pt;q=0.9,en;q=0.8',
    }
    try:
        resp = requests.get(req.url, headers=headers, timeout=15, allow_redirects=True)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, 'html.parser')

        parsed_base = urlparse(req.url)
        base_domain = parsed_base.netloc
        base_clean  = req.url.rstrip('/')

        found = set()
        for a in soup.find_all('a', href=True):
            href = a['href'].strip()
            if not href or href.startswith('#') or href.startswith('mailto:') or href.startswith('javascript:'):
                continue
            full = urljoin(req.url, href)
            p = urlparse(full)
            if p.netloc != base_domain or p.scheme not in ('http', 'https'):
                continue
            clean = f"{p.scheme}://{p.netloc}{p.path}".rstrip('/')
            if clean and clean != base_clean:
                found.add(clean)

        urls = sorted(found)[:req.max_pages]
        return {"status": "success", "base_url": req.url, "urls": urls, "total": len(urls)}

    except requests.exceptions.RequestException as e:
        return {"status": "error", "message": f"Erro ao acessar a URL: {str(e)}"}
    except Exception as e:
        print(f"[crawl_site] Erro: {e}")
        return {"status": "error", "message": str(e)}
